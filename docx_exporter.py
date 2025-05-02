from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os
from datetime import datetime

def export_to_docx(summary_text, original_audio_path):
    doc = Document()

    # 封面
    doc.add_picture("assets/qdog.png", width=Pt(200))
    title = doc.add_paragraph("Q寶筆記神器報告")
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    title.runs[0].font.size = Pt(20)
    title.runs[0].font.name = "標楷體"

    audio_name = os.path.basename(original_audio_path)
    doc.add_paragraph(f"📁 處理檔案：{audio_name}", style="Normal")
    doc.add_paragraph(f"🗓️ 產出時間：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", style="Normal")
    doc.add_paragraph("🟢 汪～任務完成，這是今天的會議紀錄！")

    doc.add_page_break()

    # 條列記錄 + 摘要區塊（由 GPT 產出）
    doc.add_heading("📋 條列式會議重點", level=1)
    doc.add_paragraph(summary_text, style="List Bullet")

    doc.add_page_break()

    # 匯出檔案名稱
    filename = os.path.splitext(audio_name)[0] + "_Q寶報告.docx"
    output_path = os.path.join("outputs", filename)
    os.makedirs("outputs", exist_ok=True)
    doc.save(output_path)

    return output_path
